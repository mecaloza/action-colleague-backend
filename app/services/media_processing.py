"""
Processing of uploaded files (pure functions over local files; the job handles storage).

Videos are probed and, when needed, re-encoded to web-friendly H.264/AAC MP4 (max 1080p, aspect
ratio preserved — never stretched) with a poster frame. Documents get their text extracted (the
AI uses it to design courses and quizzes). PDF decks are rendered to one image per page for the
recording studio. Images are resized to a sane maximum.
"""

import json
import logging
import subprocess
from dataclasses import dataclass
from pathlib import Path

from PIL import Image, ImageOps

logger = logging.getLogger(__name__)

MAX_TEXT_CHARS = 200_000
WEB_VIDEO_CODECS = {"h264"}
WEB_AUDIO_CODECS = {"aac", "mp3"}
MAX_VIDEO_WIDTH, MAX_VIDEO_HEIGHT = 1920, 1080  # 1080p: `is_web_ready` and `to_web_mp4` must agree on it
POSTER_WIDTH = 1280
PROBE_TIMEOUT_SECONDS = 120  # ffprobe and single-frame grabs
DECODE_TIMEOUT_SECONDS = 1800  # reading a whole video without encoding it
ENCODE_TIMEOUT_SECONDS = 3600

_WEB_VIDEO_FILTER = (
    # Fit inside 1080p keeping the aspect ratio (never stretch), then force even sizes for H.264.
    # The box follows the orientation: landscape fits 1920x1080, portrait fits 1080x1920.
    f"scale='if(gte(iw,ih),min({MAX_VIDEO_WIDTH},iw),min({MAX_VIDEO_HEIGHT},iw))'"
    f":'if(gte(iw,ih),min({MAX_VIDEO_HEIGHT},ih),min({MAX_VIDEO_WIDTH},ih))':force_original_aspect_ratio=decrease,"
    "scale=trunc(iw/2)*2:trunc(ih/2)*2,fps=30,format=yuv420p"
)


@dataclass
class VideoInfo:
    duration: float
    width: int
    height: int
    video_codec: str
    audio_codec: str | None
    rotation: int = 0

    @property
    def is_web_ready(self) -> bool:
        return (
            self.video_codec in WEB_VIDEO_CODECS
            and (self.audio_codec is None or self.audio_codec in WEB_AUDIO_CODECS)
            and max(self.width, self.height) <= MAX_VIDEO_WIDTH
            and min(self.width, self.height) <= MAX_VIDEO_HEIGHT
            and self.rotation == 0
        )


class MediaError(ValueError):
    """The file can't be used; the message is shown to the admin."""


def run(cmd: list[str], timeout: int = ENCODE_TIMEOUT_SECONDS) -> subprocess.CompletedProcess:
    try:
        return subprocess.run(cmd, check=True, capture_output=True, timeout=timeout)
    except subprocess.CalledProcessError as exc:
        stderr = (exc.stderr or b"")[-600:].decode(errors="replace")
        logger.warning("media_command_failed", extra={"command": cmd[0], "returncode": exc.returncode, "stderr": stderr})
        raise MediaError("No pudimos leer el archivo. ¿Está completo y es un formato de video válido?") from exc


def _first_stream(data: dict, codec_type: str) -> dict | None:
    return next((s for s in data.get("streams", []) if s.get("codec_type") == codec_type), None)


def _rotation(video_stream: dict) -> int:
    """Degrees (0-359) a player must turn the frames; phones often store this instead of turning the pixels."""
    rotation = 0
    for side in video_stream.get("side_data_list", []) or []:
        if "rotation" in side:
            rotation = abs(int(side["rotation"])) % 360
    return rotation


def probe(path: Path) -> VideoInfo:
    raw = run(
        ["ffprobe", "-v", "error", "-show_streams", "-show_format", "-of", "json", str(path)], PROBE_TIMEOUT_SECONDS
    ).stdout
    data = json.loads(raw)
    video = _first_stream(data, "video")
    if video is None:
        raise MediaError("El archivo no contiene video")
    audio = _first_stream(data, "audio")
    duration = data.get("format", {}).get("duration") or video.get("duration")
    rotation = _rotation(video)
    if duration in (None, "N/A"):
        duration = measure_duration(path)
    if not float(duration) > 0:
        raise MediaError("El video está vacío")
    return VideoInfo(
        duration=float(duration),
        width=int(video.get("width", 0)),
        height=int(video.get("height", 0)),
        video_codec=video.get("codec_name", ""),
        audio_codec=audio.get("codec_name") if audio else None,
        rotation=rotation,
    )


def measure_duration(path: Path) -> float:
    """Browser recordings (WebM) often lack a duration header: decode to find the real length."""
    result = run(["ffmpeg", "-nostdin", "-v", "error", "-progress", "pipe:1", "-i", str(path), "-map", "0:v:0",
                  "-f", "null", "-"], DECODE_TIMEOUT_SECONDS)
    seconds = 0.0
    for line in result.stdout.decode(errors="replace").splitlines():
        if line.startswith("out_time_us="):
            value = line.split("=", 1)[1]
            if value.isdigit():
                seconds = int(value) / 1_000_000
    return seconds


def to_web_mp4(src: Path, dest: Path) -> None:
    """H.264/AAC MP4 within 1920x1080, aspect ratio preserved, even dimensions, fast start."""
    run(
        [
            "ffmpeg", "-nostdin", "-y", "-v", "error", "-i", str(src),
            "-vf", _WEB_VIDEO_FILTER,
            "-c:v", "libx264", "-preset", "veryfast", "-crf", "21",
            "-c:a", "aac", "-b:a", "128k", "-ar", "48000", "-af", "aresample=async=1",
            "-movflags", "+faststart", "-threads", "2", str(dest),
        ],
        ENCODE_TIMEOUT_SECONDS,
    )


def poster(src: Path, dest: Path, at_seconds: float) -> None:
    run(["ffmpeg", "-nostdin", "-y", "-v", "error", "-ss", f"{max(at_seconds, 0):.2f}", "-i", str(src),
         "-frames:v", "1", "-vf", f"scale='min({POSTER_WIDTH},iw)':-2", "-q:v", "3", str(dest)], PROBE_TIMEOUT_SECONDS)


# ── Documents ─────────────────────────────────────────────────────────


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    slides = []
    for index, slide in enumerate(Presentation(str(path)).slides, start=1):
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        slides.append(f"[Diapositiva {index}]\n" + "\n".join(t for t in texts if t.strip()))
    return "\n\n".join(slides)


def extract_text(path: Path, mime_type: str, filename: str) -> str:
    # Each reader imports its library on use, so importing this module stays cheap.
    name = filename.lower()
    if mime_type == "application/pdf" or name.endswith(".pdf"):
        text = _pdf_text(path)
    elif name.endswith(".docx"):
        text = _docx_text(path)
    elif name.endswith(".pptx"):
        text = _pptx_text(path)
    elif mime_type.startswith("text/") or name.endswith((".txt", ".md", ".csv")):
        text = path.read_text(errors="replace")
    else:
        raise MediaError("Formato de documento no soportado (usa PDF, DOCX, PPTX, TXT o MD)")
    return text.strip()[:MAX_TEXT_CHARS]


def extract_text_safely(path: Path, mime_type: str, filename: str) -> str:
    """extract_text, with any parser failure reported as an unreadable document."""
    try:
        return extract_text(path, mime_type, filename)
    except MediaError:
        raise
    except Exception as exc:  # pypdf, python-docx and python-pptx raise many kinds of errors
        logger.warning("document_unreadable", extra={"error_type": type(exc).__name__, "error": str(exc)[:300]})
        raise MediaError("No pudimos leer el documento. ¿Está protegido con contraseña o dañado?") from exc


def render_pdf_pages(path: Path, out_dir: Path, width: int = 1920, max_pages: int = 150) -> list[Path]:
    """One PNG per page (for presenting slides while recording)."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        if len(pdf) > max_pages:
            raise MediaError(f"La presentación tiene {len(pdf)} páginas; el máximo es {max_pages}")
        pages = []
        for index in range(len(pdf)):
            page = pdf[index]
            scale = width / page.get_width()
            image = page.render(scale=scale).to_pil().convert("RGB")
            target = out_dir / f"{index + 1:03d}.png"
            image.save(target, optimize=True)
            pages.append(target)
        return pages
    finally:
        pdf.close()


def image_size(path: Path) -> tuple[int, int]:
    """(width, height) of an image file."""
    with Image.open(path) as image:
        return image.size


def normalize_image(src: Path, dest: Path, max_side: int = 1920) -> tuple[int, int]:
    """JPEG within max_side x max_side, EXIF orientation applied. Returns (width, height)."""
    try:
        with Image.open(src) as image:
            image = ImageOps.exif_transpose(image).convert("RGB")
            image.thumbnail((max_side, max_side))
            image.save(dest, "JPEG", quality=88, optimize=True)
            return image.size
    except OSError as exc:
        raise MediaError("La imagen no se pudo leer (usa PNG, JPG o WEBP)") from exc
